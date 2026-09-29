"""
DOCX Document Renderer for JobPilot AI.
Generates structured Microsoft Word documents using python-docx for portals explicitly mandating Word format.
"""

from pathlib import Path
from typing import Optional, Union
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from jobpilot.core.registry import PluginRegistry
from jobpilot.core.schemas import CoverLetterDoc, ResumeDoc
from jobpilot.documents.pdf_renderer import get_default_cover_letters_dir, get_default_resumes_dir


@PluginRegistry.register_renderer("docx")
class DocxRenderer:
    """Renders ResumeDoc and CoverLetterDoc into structured .docx documents."""

    def __init__(self, output_dir: Optional[Path] = None):
        self.output_dir = Path(output_dir) if output_dir else None

    def render_resume(
        self,
        resume: ResumeDoc,
        output_path: Optional[Union[str, Path]] = None,
    ) -> Path:
        """Renders ResumeDoc to a .docx file."""
        if not output_path:
            out_dir = self.output_dir or get_default_resumes_dir()
            filename = f"Resume_{resume.full_name.replace(' ', '_')}_{resume.variant}.docx"
            target_path = out_dir / filename
        else:
            target_path = Path(output_path)

        target_path.parent.mkdir(parents=True, exist_ok=True)

        doc = docx.Document()

        # Set 0.5-inch margins
        for section in doc.sections:
            section.top_margin = Inches(0.5)
            section.bottom_margin = Inches(0.5)
            section.left_margin = Inches(0.5)
            section.right_margin = Inches(0.5)

        # Name Title
        title_p = doc.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title_p.add_run(resume.full_name.upper())
        title_run.font.name = "Arial"
        title_run.font.size = Pt(20)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(26, 54, 93)

        # Contact line
        contact_items = [resume.location]
        if resume.contact_email and resume.contact_email != "NEEDS_USER_INPUT":
            contact_items.append(resume.contact_email)
        if resume.contact_phone and resume.contact_phone != "NEEDS_USER_INPUT":
            contact_items.append(resume.contact_phone)
        contact_p = doc.add_paragraph(" | ".join(contact_items))
        contact_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        contact_p.runs[0].font.size = Pt(9.5)
        contact_p.runs[0].font.color.rgb = RGBColor(100, 100, 100)

        # Section: Summary
        doc.add_heading("Professional Summary", level=1)
        doc.add_paragraph(resume.summary)

        # Section: Skills
        doc.add_heading("Technical Competencies", level=1)
        for cat, items in resume.skills.items():
            if items:
                p = doc.add_paragraph()
                p.add_run(f"{cat.replace('_', ' ').title()}: ").bold = True
                p.add_run(", ".join(items))

        # Section: Experience
        if resume.experience:
            doc.add_heading("Technical Experience", level=1)
            for exp in resume.experience:
                employer = exp.get("employer", "")
                title = exp.get("title", "")
                if employer or title:
                    role_p = doc.add_paragraph()
                    role_p.add_run(f"{title} — {employer}").bold = True
                    for b in exp.get("bullets", []):
                        doc.add_paragraph(b, style="List Bullet")

        # Section: Projects
        if resume.projects:
            doc.add_heading("Featured Projects", level=1)
            for proj in resume.projects:
                name = proj.get("name", "")
                if name:
                    proj_p = doc.add_paragraph()
                    proj_p.add_run(name).bold = True
                    for b in proj.get("bullets", []):
                        doc.add_paragraph(b, style="List Bullet")

        # Section: Education
        doc.add_heading("Education", level=1)
        for edu in resume.education:
            deg = edu.get("degree", "")
            field = edu.get("field", "")
            inst = edu.get("institution", "")
            yr = edu.get("graduation_year", "")
            p = doc.add_paragraph()
            p.add_run(f"{deg} in {field}").bold = True
            sub = []
            if inst and inst != "NEEDS_USER_INPUT":
                sub.append(inst)
            if yr and yr != "NEEDS_USER_INPUT":
                sub.append(str(yr))
            if sub:
                p.add_run(f" ({', '.join(sub)})").italic = True

        doc.save(str(target_path))
        return target_path

    def render_cover_letter(
        self,
        cover_letter: CoverLetterDoc,
        output_path: Optional[Union[str, Path]] = None,
    ) -> Path:
        """Renders CoverLetterDoc to a .docx file."""
        if not output_path:
            out_dir = self.output_dir or get_default_cover_letters_dir()
            filename = f"CoverLetter_{cover_letter.company.replace(' ', '_')}_{cover_letter.job_id}.docx"
            target_path = out_dir / filename
        else:
            target_path = Path(output_path)

        target_path.parent.mkdir(parents=True, exist_ok=True)

        doc = docx.Document()

        # Name Title
        name_p = doc.add_paragraph()
        name_run = name_p.add_run(cover_letter.applicant_name)
        name_run.font.name = "Arial"
        name_run.font.size = Pt(18)
        name_run.font.bold = True
        name_run.font.color.rgb = RGBColor(26, 54, 93)

        doc.add_paragraph(cover_letter.salutation).bold = True
        doc.add_paragraph(cover_letter.opening)

        for p in cover_letter.body_paragraphs:
            doc.add_paragraph(p)

        doc.add_paragraph(cover_letter.closing)

        sign_off_p = doc.add_paragraph(f"\n{cover_letter.sign_off}\n\n")
        sign_off_p.add_run(cover_letter.applicant_name).bold = True

        doc.save(str(target_path))
        return target_path
