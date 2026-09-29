"""
Tests for JobPilot AI Document Generation Layer.
Verifies format detection, HTML compilation, Playwright PDF rendering, and DOCX generation.
"""

from pathlib import Path
import pytest
import docx
from jobpilot.applications.cover_letter_gen import CoverLetterGenerator
from jobpilot.applications.resume_master import ResumeMaster
from jobpilot.core.registry import PluginRegistry
from jobpilot.core.schemas import DocumentFormat, JobRecord
from jobpilot.documents.docx_renderer import DocxRenderer
from jobpilot.documents.format_detector import FormatDetector
from jobpilot.documents.html_template import render_cover_letter_html, render_resume_html
from jobpilot.documents.pdf_renderer import PDFRenderer


def create_sample_job() -> JobRecord:
    return JobRecord(
        job_id="job_doc_test_01",
        company="OpenAI",
        title="Solutions Engineer - Enterprise AI",
        location="Remote",
        country="India",
        description="Looking for an AI engineer to integrate enterprise models. Please submit your application in PDF or Word format.",
        skills=["Python", "REST APIs", "LLMs"],
        job_url="https://openai.com/careers/1",
        source="test",
        date_found="2026-09-29",
    )


class TestFormatDetector:
    def test_default_is_pdf(self):
        job = create_sample_job()
        fmt = FormatDetector.detect_format(job=job)
        assert fmt == DocumentFormat.PDF

    def test_detect_explicit_word_requirement_in_accept_attr(self):
        fmt = FormatDetector.detect_format(portal_accept_attr=".doc,.docx,application/msword")
        assert fmt == DocumentFormat.DOCX

    def test_detect_explicit_docx_in_job_description(self):
        job = JobRecord(
            job_id="j2",
            company="LegacyCorp",
            title="Systems Dev",
            description="All candidates must upload resume in Word format (.docx only).",
            job_url="https://example.com",
            source="test",
            date_found="2026-09-29",
        )
        fmt = FormatDetector.detect_format(job=job)
        assert fmt == DocumentFormat.DOCX

    def test_detect_txt_format(self):
        fmt = FormatDetector.detect_format(portal_accept_attr=".txt,text/plain")
        assert fmt == DocumentFormat.TXT


class TestHTMLTemplates:
    def test_render_resume_html(self):
        master = ResumeMaster()
        resume = master.assemble_resume()
        html = render_resume_html(resume)

        assert "<html" in html
        assert "Ankita Yadav" in html
        assert "Technical Competencies" in html
        assert "Python" in html
        assert "Education" in html

    def test_render_cover_letter_html(self):
        job = create_sample_job()
        gen = CoverLetterGenerator()
        cl = gen.generate(job)
        html = render_cover_letter_html(cl)

        assert "<html" in html
        assert "OpenAI" in html
        assert "Ankita Yadav" in html
        assert "Dear OpenAI Hiring Team," in html


class TestDocumentRenderers:
    def test_plugin_registry_renderers(self):
        pdf_cls = PluginRegistry.get_renderer("pdf")
        docx_cls = PluginRegistry.get_renderer("docx")
        assert pdf_cls is PDFRenderer
        assert docx_cls is DocxRenderer

    def test_render_pdf_resume_and_cover_letter(self, tmp_path):
        master = ResumeMaster()
        resume = master.assemble_resume()
        job = create_sample_job()
        cl = CoverLetterGenerator().generate(job)

        renderer = PDFRenderer(output_dir=tmp_path)

        # 1. Resume PDF
        pdf_resume_path = renderer.render_resume(resume, output_path=tmp_path / "resume_test.pdf")
        assert pdf_resume_path.exists()
        assert pdf_resume_path.stat().st_size > 1000
        # Check standard PDF header
        with open(pdf_resume_path, "rb") as f:
            header = f.read(5)
            assert header.startswith(b"%PDF")

        # 2. Cover Letter PDF
        pdf_cl_path = renderer.render_cover_letter(cl, output_path=tmp_path / "cover_letter_test.pdf")
        assert pdf_cl_path.exists()
        assert pdf_cl_path.stat().st_size > 1000
        with open(pdf_cl_path, "rb") as f:
            header = f.read(5)
            assert header.startswith(b"%PDF")

    def test_render_docx_resume_and_cover_letter(self, tmp_path):
        master = ResumeMaster()
        resume = master.assemble_resume()
        job = create_sample_job()
        cl = CoverLetterGenerator().generate(job)

        renderer = DocxRenderer(output_dir=tmp_path)

        # 1. Resume DOCX
        docx_resume_path = renderer.render_resume(resume, output_path=tmp_path / "resume_test.docx")
        assert docx_resume_path.exists()
        assert docx_resume_path.stat().st_size > 1000
        # Verify it can be read back by python-docx
        doc = docx.Document(str(docx_resume_path))
        full_text = " ".join([p.text for p in doc.paragraphs])
        assert "ANKITA YADAV" in full_text
        assert "Technical Competencies" in full_text

        # 2. Cover Letter DOCX
        docx_cl_path = renderer.render_cover_letter(cl, output_path=tmp_path / "cl_test.docx")
        assert docx_cl_path.exists()
        assert docx_cl_path.stat().st_size > 1000
        doc_cl = docx.Document(str(docx_cl_path))
        cl_text = " ".join([p.text for p in doc_cl.paragraphs])
        assert "OpenAI" in cl_text
        assert "Ankita Yadav" in cl_text
